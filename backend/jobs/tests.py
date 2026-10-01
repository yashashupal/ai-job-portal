from io import BytesIO
from unittest.mock import MagicMock, patch

from docx import Document
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.cache import cache
from django.test import override_settings
from rest_framework_simplejwt.tokens import RefreshToken
from rest_framework.test import APITestCase

from accounts.models import User
from jobs.models import ExternalJobEmbedding


def register(client, email, role, **extra):
    include_resume = extra.pop("include_resume", role == "applicant")
    body = {"email": email, "password": "StrongPass123", "first_name": "Test", "role": role, **extra}
    if include_resume:
        document = Document()
        document.add_paragraph("Python React full stack engineer with API, database, and production experience. " * 5)
        buffer = BytesIO()
        document.save(buffer)
        body["resume"] = SimpleUploadedFile("resume.docx", buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        r = client.post("/api/auth/register/", body, format="multipart")
    else:
        r = client.post("/api/auth/register/", body, format="json")
    assert r.status_code == 201, r.content
    return r.data["access"]


class PortalFlowTests(APITestCase):
    def auth(self, token):
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {token}")

    def test_full_flow(self):
        hirer = register(self.client, "boss@acme.com", "hirer", company_name="Acme")
        self.auth(hirer)
        r = self.client.post("/api/jobs/", {
            "title": "Django Developer", "location": "Indore", "work_mode": "hybrid",
            "job_type": "full_time", "description": "Build APIs.", "skills": ["Python", "python", "Django"],
            "salary_min": 600000, "salary_max": 900000,
        }, format="json")
        self.assertEqual(r.status_code, 201, r.content)
        job_id = r.data["id"]
        self.assertEqual(r.data["company_name"], "Acme")
        self.assertEqual(r.data["skills"], ["Python", "Django"])

        self.client.credentials()
        r = self.client.get("/api/jobs/?q=django")
        self.assertEqual(r.data["count"], 1)

        applicant = register(self.client, "dev@mail.com", "applicant")
        self.auth(applicant)
        r = self.client.post("/api/jobs/", {"title": "x"}, format="json")
        self.assertEqual(r.status_code, 403)

        resume = SimpleUploadedFile("cv.pdf", b"%PDF-1.4 test", content_type="application/pdf")
        if True:
            r = self.client.post(f"/api/jobs/{job_id}/apply/", {"cover_letter": "Hi", "resume": resume}, format="multipart")
            self.assertEqual(r.status_code, 201, r.content)
            resume2 = SimpleUploadedFile("cv.pdf", b"%PDF-1.4 test", content_type="application/pdf")
            r = self.client.post(f"/api/jobs/{job_id}/apply/", {"resume": resume2}, format="multipart")
            self.assertEqual(r.status_code, 400)
            bad = SimpleUploadedFile("cv.pdf", b"MZ not a pdf", content_type="application/octet-stream")
            r = self.client.post(f"/api/jobs/{job_id}/apply/", {"resume": bad}, format="multipart")
            self.assertEqual(r.status_code, 400)

        r = self.client.get("/api/applications/mine/")
        self.assertEqual(len(r.data), 1)
        self.assertEqual(r.data[0]["job"]["title"], "Django Developer")
        r = self.client.get(f"/api/jobs/{job_id}/")
        self.assertTrue(r.data["has_applied"])

        # applicant cannot see applicants
        self.assertEqual(self.client.get(f"/api/jobs/{job_id}/applications/").status_code, 403)

        self.auth(hirer)
        r = self.client.get(f"/api/jobs/{job_id}/applications/")
        self.assertEqual(len(r.data), 1)
        app_id = r.data[0]["id"]
        d = self.client.get(f"/api/applications/{app_id}/resume/")
        self.assertEqual(d.status_code, 200)
        self.assertEqual(d["Content-Type"], "application/pdf")
        self.assertTrue(d.content.startswith(b"%PDF"))
        r = self.client.patch(f"/api/applications/{app_id}/", {"status": "shortlisted"}, format="json")
        self.assertEqual(r.data["status"], "shortlisted")
        self.assertEqual(self.client.get("/api/jobs/mine/").data[0]["applications_count"], 1)

        # another hirer cannot touch it
        other = register(self.client, "other@x.com", "hirer", company_name="X")
        self.auth(other)
        self.assertEqual(self.client.get(f"/api/applications/{app_id}/resume/").status_code, 404)
        self.client.credentials()
        self.assertEqual(self.client.get(f"/api/applications/{app_id}/resume/").status_code, 401)
        self.auth(other)
        self.assertEqual(self.client.delete(f"/api/jobs/{job_id}/").status_code, 404)
        self.assertEqual(self.client.patch(f"/api/applications/{app_id}/", {"status": "hired"}, format="json").status_code, 404)


class RegistrationResumeTests(APITestCase):
    def test_applicant_registration_requires_and_saves_resume(self):
        body = {"email": "new-applicant@example.com", "password": "StrongPass123", "first_name": "New", "role": "applicant"}
        missing = self.client.post("/api/auth/register/", body, format="json")
        self.assertEqual(missing.status_code, 400)
        self.assertIn("resume", missing.data)

        document = Document()
        document.add_paragraph("Full stack engineer skilled in Python, React, PostgreSQL, and REST API development. " * 5)
        buffer = BytesIO()
        document.save(buffer)
        body["resume"] = SimpleUploadedFile("new-resume.docx", buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        created = self.client.post("/api/auth/register/", body, format="multipart")
        self.assertEqual(created.status_code, 201, created.content)
        self.assertTrue(created.data["user"]["has_resume"])
        applicant = User.objects.get(email="new-applicant@example.com")
        self.assertIn("PostgreSQL", applicant.resume_text)

    def test_hirer_registration_does_not_require_resume(self):
        response = self.client.post("/api/auth/register/", {
            "email": "new-hirer@example.com", "password": "StrongPass123", "first_name": "New",
            "role": "hirer", "company_name": "Acme",
        }, format="json")
        self.assertEqual(response.status_code, 201, response.content)

    def test_login(self):
        register(self.client, "a@b.com", "applicant")
        self.assertEqual(self.client.post("/api/auth/login/", {"email": "A@B.com", "password": "StrongPass123"}, format="json").status_code, 200)
        self.assertEqual(self.client.post("/api/auth/login/", {"email": "a@b.com", "password": "nope"}, format="json").status_code, 401)

    def test_hirer_needs_company(self):
        r = self.client.post("/api/auth/register/", {"email": "h@h.com", "password": "StrongPass123", "role": "hirer"}, format="json")
        self.assertEqual(r.status_code, 400)


FAKE = {
    "data": [{
        "id": "1", "job_title": "React Engineer", "company": "Notion", "company_domain": "notion.so",
        "location": "Remote", "remote": True, "min_annual_salary_usd": 120000, "max_annual_salary_usd": 160000,
        "technology_slugs": ["react"], "description": "## Build things",
        "sources": [{"provider": "greenhouse", "url": "https://x.co/1"}],
    }],
    "metadata": {"next_cursor": "abc", "total": 42},
}


class ExternalTests(APITestCase):
    def setUp(self):
        cache.clear()

    @override_settings(JOBSPIPE_API_KEY="jp_live_test")
    @patch("jobs.jobspipe.requests.post")
    def test_search_uses_bearer_and_normalizes(self, post):
        post.return_value = MagicMock(ok=True, status_code=200, json=lambda: FAKE)
        r = self.client.get("/api/external/jobs/?q=react&remote=1")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data["results"][0]["apply_url"], "https://x.co/1")
        self.assertEqual(r.data["next_cursor"], "abc")
        _, kwargs = post.call_args
        self.assertEqual(kwargs["headers"]["Authorization"], "Bearer jp_live_test")
        self.assertEqual(kwargs["json"]["job_title_or"], ["react"])
        self.assertNotIn("jp_live_test", str(r.content))

    @override_settings(JOBSPIPE_API_KEY="jp_live_test")
    @patch("jobs.jobspipe.requests.post")
    def test_stack(self, post):
        post.return_value = MagicMock(ok=True, status_code=200, json=lambda: {"technologies": [
            {"name": "Next.js", "category": "Framework", "confidence": 90}, {"name": "AWS", "confidence": 60}]})
        r = self.client.get("/api/external/stack/?domain=https://www.notion.so/path")
        self.assertEqual(r.data["domain"], "notion.so")
        self.assertEqual(r.data["technologies"][0]["name"], "Next.js")
        self.assertEqual(self.client.get("/api/external/stack/?domain=not_a_domain").status_code, 400)

    @override_settings(JOBSPIPE_API_KEY="jp_live_bad")
    @patch("jobs.jobspipe.requests.post")
    def test_bad_key(self, post):
        post.return_value = MagicMock(ok=False, status_code=401)
        self.assertEqual(self.client.get("/api/external/jobs/").status_code, 502)


class ApplicantProfileTests(APITestCase):
    def test_profile_edit_and_private_resume_lifecycle(self):
        token = register(self.client, "profile@example.com", "applicant")
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {token}")

        updated = self.client.patch("/api/auth/me/", {"headline": "Backend engineer", "first_name": "Asha"}, format="json")
        self.assertEqual(updated.status_code, 200)
        self.assertEqual(updated.data["headline"], "Backend engineer")

        document = Document()
        document.add_paragraph("Python Django PostgreSQL REST API engineer with production experience. " * 8)
        buffer = BytesIO()
        document.save(buffer)
        upload = SimpleUploadedFile("asha-resume.docx", buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        response = self.client.post("/api/auth/me/resume/", {"resume": upload}, format="multipart")
        self.assertEqual(response.status_code, 200, response.content)
        self.assertTrue(response.data["has_resume"])

        user = User.objects.get(email="profile@example.com")
        self.assertIn("Django", user.resume_text)
        self.assertEqual(user.resume_embedding, [])
        downloaded = self.client.get("/api/auth/me/resume/")
        self.assertEqual(downloaded.status_code, 200)
        self.assertEqual(downloaded.content, bytes(user.resume_data))
        self.assertEqual(self.client.delete("/api/auth/me/resume/").status_code, 204)
        user.refresh_from_db()
        self.assertIsNone(user.resume_data)
        self.assertEqual(user.resume_text, "")

    def test_resume_endpoint_requires_authentication(self):
        self.assertEqual(self.client.get("/api/auth/me/resume/").status_code, 401)


class ExternalTrackingTests(APITestCase):
    def test_saved_live_jobs_are_persistent_and_private_to_the_owner(self):
        token = register(self.client, "tracker@example.com", "applicant")
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {token}")
        job = {"title": "Python Engineer", "company": "Acme", "apply_url": "https://jobs.example.com/1", "technologies": ["Python"]}
        created = self.client.post("/api/external/tracked/", {"external_id": "source-1", "job": job}, format="json")
        self.assertEqual(created.status_code, 200, created.content)
        self.assertEqual(created.data["status"], "saved")
        self.assertEqual(len(self.client.get("/api/external/tracked/").data), 1)

        updated = self.client.patch(f"/api/external/tracked/{created.data['id']}/", {"status": "applied"}, format="json")
        self.assertEqual(updated.data["status"], "applied")

        other = register(self.client, "other-tracker@example.com", "applicant")
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {other}")
        self.assertEqual(self.client.patch(f"/api/external/tracked/{created.data['id']}/", {"status": "saved"}, format="json").status_code, 404)
        self.assertEqual(self.client.delete(f"/api/external/tracked/{created.data['id']}/").status_code, 404)

    def test_invalid_external_apply_urls_are_rejected(self):
        token = register(self.client, "url-check@example.com", "applicant")
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {token}")
        response = self.client.post("/api/external/tracked/", {
            "external_id": "source-2", "job": {"title": "Python Engineer", "apply_url": "javascript:alert(1)"},
        }, format="json")
        self.assertEqual(response.status_code, 400)


class GeminiMatchingTests(APITestCase):
    @override_settings(GEMINI_API_KEY="test-key", GEMINI_EMBEDDING_MODEL="gemini-embedding-001", GEMINI_TEXT_MODEL="gemini-2.5-flash")
    @patch("jobs.gemini._client", return_value=object())
    @patch("jobs.gemini._match_insights", side_effect=lambda client, resume, jobs: {"summary": "Matches ranked.", "results": jobs})
    @patch("jobs.gemini._embed", side_effect=[[[1.0, 0.0]], [[0.0, 1.0], [1.0, 0.0]]])
    def test_resume_and_job_embeddings_rank_and_persist(self, embed, insights, client):
        from jobs.gemini import rank_external_jobs

        user = User.objects.create_user(username="match@example.com", email="match@example.com", password="StrongPass123")
        user.resume_text = "Python Django backend engineer with PostgreSQL and REST experience." * 4
        user.save(update_fields=("resume_text",))
        jobs = [
            {"id": "one", "title": "Frontend Artist", "technologies": ["CSS"]},
            {"id": "two", "title": "Python Engineer", "technologies": ["Django"]},
        ]

        result = rank_external_jobs(user, jobs)
        self.assertEqual([job["id"] for job in result["results"]], ["two", "one"])
        self.assertEqual(result["results"][0]["match_score"], 100)
        self.assertEqual(ExternalJobEmbedding.objects.count(), 2)
        self.assertTrue(User.objects.get(pk=user.pk).resume_embedding)

        rank_external_jobs(User.objects.get(pk=user.pk), jobs)
        self.assertEqual(embed.call_count, 2)

    @patch("jobs.views.jobspipe.search_jobs")
    def test_recommendations_do_not_consume_search_without_a_resume(self, search):
        user = User.objects.create_user(username="no-resume@example.com", email="no-resume@example.com", password="StrongPass123")
        token = str(RefreshToken.for_user(user).access_token)
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {token}")
        response = self.client.get("/api/external/jobs/recommended/")
        self.assertEqual(response.status_code, 400)
        self.assertIn("Upload your resume", response.data["detail"])
        search.assert_not_called()

    def test_recommendations_search_by_headline_and_limit_initial_batch(self):
        token = register(self.client, "headline@example.com", "applicant")
        user = User.objects.get(email="headline@example.com")
        user.headline = "Full stack developer"
        user.resume_text = "React Python full stack developer experience." * 5
        user.save(update_fields=("headline", "resume_text"))
        self.client.credentials(HTTP_AUTHORIZATION=f"Bearer {token}")
        search_result = {"results": [], "next_cursor": None, "total": 0, "sandbox": False}
        with override_settings(GEMINI_API_KEY="test-key"), \
             patch("jobs.views.jobspipe.search_jobs", return_value=search_result) as search, \
             patch("jobs.views.gemini.rank_external_jobs", return_value={"summary": "Ranked.", "results": []}):
            response = self.client.get("/api/external/jobs/recommended/")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(search.call_args.kwargs["q"], "Full stack developer")
        self.assertEqual(search.call_args.kwargs["limit"], 10)

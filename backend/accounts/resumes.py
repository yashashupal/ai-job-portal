from io import BytesIO

from docx import Document
from pypdf import PdfReader


class ResumeTextError(ValueError):
    pass


def extract_resume_text(data, filename):
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    try:
        if suffix == "pdf":
            reader = PdfReader(BytesIO(data))
            text = "\n".join(page.extract_text() or "" for page in reader.pages[:30])
        elif suffix == "docx":
            document = Document(BytesIO(data))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            paragraphs.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            text = "\n".join(paragraphs)
        else:
            raise ResumeTextError("Upload a PDF or DOCX resume.")
    except ResumeTextError:
        raise
    except Exception as exc:
        raise ResumeTextError("We could not read this resume. Try another PDF or DOCX file.") from exc

    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())[:30000]
    if len(text) < 80:
        raise ResumeTextError("We could not find readable text. Upload a text-based PDF or DOCX resume.")
    return text